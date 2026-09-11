"""
PrivacyGuard
Small Flask app for text-only PII detection and redaction.
"""

from __future__ import annotations

import atexit
import logging
import re
import shutil
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import PyPDF2
from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor as DocxRGBColor
from flask import Flask, jsonify, render_template, request, send_file, session
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from werkzeug.utils import secure_filename

try:
    import fitz
except ImportError:
    fitz = None

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)
app.config["SECRET_KEY"] = "privacyguard-dev-secret-change-me"
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)

SESSION_ROOT = Path("session_files")
SESSION_ROOT.mkdir(exist_ok=True)

SUPPORTED_EXTENSIONS = {
    ".txt": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
}

SESSION_DOCS: Dict[str, Dict[str, Dict]] = {}
SESSION_TOUCHED: Dict[str, datetime] = {}
SESSION_TTL = timedelta(hours=8)


PII_PATTERNS = [
    {
        "type": "US_SSN",
        "label": "Social Security Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?!000|666|9\d{2})\d{3}[- ]?(?!00)\d{2}[- ]?(?!0000)\d{4}\b"),
    },
    {
        "type": "CREDIT_CARD",
        "label": "Credit Card Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:\d[ -]*?){13,19}\b"),
        "validator": "luhn",
    },
    {
        "type": "CARD_EXPIRATION",
        "label": "Card Expiration Date",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b(?:exp(?:ires|iration)?|valid\s+thru|good\s+thru)?\s*(?P<value>(?:0[1-9]|1[0-2])\s*[/.-]\s*(?:\d{2}|\d{4}))\b", re.I),
    },
    {
        "type": "CARD_SECURITY_CODE",
        "label": "Card Security Code",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:cvv|cvc|cid|security\s+code)\s*[:#=-]?\s*(?P<value>\d{3,4})\b", re.I),
    },
    {
        "type": "US_DRIVER_LICENSE",
        "label": "Driver License Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:dl|dln|driver'?s?\s+license|license|lic)\s*(?:no\.?|number|#|id)?\s*[:#=-]?\s*(?P<value>[A-Z0-9-]{5,20})\b", re.I),
    },
    {
        "type": "STATE_ID",
        "label": "State ID Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:state\s+id|identification\s+card|id\s*(?:no\.?|number|#))\s*[:#=-]?\s*(?P<value>[A-Z0-9-]{5,20})\b", re.I),
    },
    {
        "type": "US_PASSPORT",
        "label": "Passport Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:passport\s*(?:no\.?|number|#)?\s*[:#=-]?\s*)?(?P<value>[A-Z][0-9]{7,9})\b", re.I),
    },
    {
        "type": "US_ROUTING_NUMBER",
        "label": "Bank Routing Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:routing|aba|rtn|ach)\s*(?:number|no\.?|#)?\s*[:#=-]?\s*(?P<value>\d{9})\b", re.I),
    },
    {
        "type": "BANK_ACCOUNT",
        "label": "Bank Account Number",
        "risk": "HIGH",
        "regex": re.compile(r"\b(?:account|acct)\s*(?:number|no\.?|#)?\s*[:#=-]?\s*(?P<value>[Xx*\d-]{6,20})\b", re.I),
    },
    {
        "type": "EMAIL_ADDRESS",
        "label": "Email Address",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I),
    },
    {
        "type": "PHONE_NUMBER",
        "label": "Phone Number",
        "risk": "MEDIUM",
        "regex": re.compile(r"(?<!\w)(?:\+?1[\s.-]?)?(?:\(?\d{3}\)?[\s.-]?)\d{3}[\s.-]?\d{4}(?!\w)"),
    },
    {
        "type": "DATE_OF_BIRTH",
        "label": "Date of Birth",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b(?:dob|date\s+of\s+birth|birth\s+date)\s*[:#=-]?\s*(?P<value>\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4})\b", re.I),
    },
    {
        "type": "PASSWORD",
        "label": "Password or Secret",
        "risk": "CRITICAL",
        "regex": re.compile(r"\b(?:password|passwd|pwd|passphrase|secret)\s*[:#=]\s*(?P<value>\S{4,})\b", re.I),
    },
    {
        "type": "API_KEY",
        "label": "API Key or Token",
        "risk": "CRITICAL",
        "regex": re.compile(r"\b(?:api[_\s-]?key|access[_\s-]?token|auth[_\s-]?token|bearer)\s*[:#=]?\s*(?P<value>[A-Za-z0-9._~+/=-]{16,})\b", re.I),
    },
    {
        "type": "USERNAME",
        "label": "Username",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b(?:username|user\s*name|login\s*id|user\s*id)\s*[:#=]\s*(?P<value>[A-Za-z0-9_.@-]{3,64})\b", re.I),
    },
    {
        "type": "IP_ADDRESS",
        "label": "IP Address",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b(?:(?:25[0-5]|2[0-4]\d|1?\d?\d)\.){3}(?:25[0-5]|2[0-4]\d|1?\d?\d)\b"),
    },
    {
        "type": "ADDRESS",
        "label": "Street Address",
        "risk": "MEDIUM",
        "regex": re.compile(r"\b\d{1,6}\s+[A-Z][A-Za-z0-9.' -]{2,}\s+(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln|Boulevard|Blvd|Court|Ct|Way)\b", re.I),
    },
]

RISK_SCORE = {"LOW": 1, "MEDIUM": 2, "HIGH": 3, "CRITICAL": 4}
MASK_CHAR = "█"


def get_session_id() -> str:
    if "sid" not in session:
        session["sid"] = str(uuid.uuid4())
        session.permanent = True
    sid = session["sid"]
    SESSION_TOUCHED[sid] = datetime.now()
    SESSION_DOCS.setdefault(sid, {})
    (SESSION_ROOT / sid).mkdir(parents=True, exist_ok=True)
    return sid


@app.before_request
def prepare_session() -> None:
    cleanup_expired_sessions()
    get_session_id()


def cleanup_expired_sessions() -> None:
    now = datetime.now()
    expired = [sid for sid, touched in SESSION_TOUCHED.items() if now - touched > SESSION_TTL]
    for sid in expired:
        cleanup_session(sid)


def cleanup_session(sid: str) -> None:
    SESSION_DOCS.pop(sid, None)
    SESSION_TOUCHED.pop(sid, None)
    shutil.rmtree(SESSION_ROOT / sid, ignore_errors=True)


def cleanup_all_sessions() -> None:
    SESSION_DOCS.clear()
    SESSION_TOUCHED.clear()
    shutil.rmtree(SESSION_ROOT, ignore_errors=True)


atexit.register(cleanup_all_sessions)


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload TXT, PDF, DOCX, or PPTX.")
    return SUPPORTED_EXTENSIONS[ext]


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: Path, file_type: str) -> Tuple[str, Dict]:
    if file_type == "text":
        text = read_text_file(path)
        return text, {"extraction_method": "Plain text"}

    if file_type == "pdf":
        parts = []
        page_text = []
        with path.open("rb") as handle:
            reader = PyPDF2.PdfReader(handle)
            page_count = len(reader.pages)
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                page_text.append({"page": page_num, "text": text})
                if text.strip():
                    parts.append(text)
        text = "\n\n".join(parts)
        return text, {
            "extraction_method": "PDF embedded text",
            "pages": page_count,
            "page_text": page_text,
        }

    if file_type == "docx":
        doc = DocxDocument(path)
        parts = [para.text for para in doc.paragraphs if para.text]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts), {
            "extraction_method": "DOCX direct text",
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
        }

    if file_type == "pptx":
        prs = Presentation(path)
        slides = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            shape_text = [
                shape.text for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if shape_text:
                slides.append(f"[SLIDE {slide_num}]\n" + "\n".join(shape_text))
        return "\n\n".join(slides), {
            "extraction_method": "PPTX direct text",
            "slides": len(prs.slides),
        }

    raise ValueError("Unsupported file type.")


def luhn_valid(value: str) -> bool:
    digits = [int(ch) for ch in re.sub(r"\D", "", value)]
    if not 13 <= len(digits) <= 19:
        return False
    checksum = 0
    parity = len(digits) % 2
    for idx, digit in enumerate(digits):
        if idx % 2 == parity:
            digit *= 2
            if digit > 9:
                digit -= 9
        checksum += digit
    return checksum % 10 == 0


def mask_value(value: str) -> str:
    clean = value.strip()
    if len(clean) <= 4:
        return "*" * len(clean)
    return f"{clean[:1]}{'*' * min(len(clean) - 2, 10)}{clean[-1:]}"


def match_value_span(match: re.Match) -> Tuple[int, int, str]:
    if "value" in match.re.groupindex and match.group("value"):
        return match.start("value"), match.end("value"), match.group("value").strip()
    return match.start(), match.end(), match.group(0).strip()


def scan_text(text: str, base_offset: int = 0, page: int | None = None) -> List[Dict]:
    findings = []
    seen = set()
    for pattern in PII_PATTERNS:
        for match in pattern["regex"].finditer(text):
            start, end, value = match_value_span(match)
            if pattern.get("validator") == "luhn" and not luhn_valid(value):
                continue
            absolute_start = base_offset + start
            absolute_end = base_offset + end
            key = (pattern["type"], absolute_start, absolute_end, value.lower(), page)
            if key in seen:
                continue
            seen.add(key)
            finding = {
                "type": pattern["type"],
                "display_name": pattern["label"],
                "value": mask_value(value),
                "raw_value": value,
                "start": absolute_start,
                "end": absolute_end,
                "length": absolute_end - absolute_start,
                "risk_level": pattern["risk"],
                "detection_method": "regex",
            }
            if page is not None:
                finding["page"] = page
            findings.append(finding)
    return assign_finding_ids(merge_overlaps(findings))


def scan_pdf_pages(page_text: List[Dict]) -> List[Dict]:
    findings = []
    offset = 0
    for page in page_text:
        text = page.get("text", "")
        page_number = page.get("page")
        findings.extend(scan_text(text, base_offset=offset, page=page_number))
        offset += len(text) + 2
    return assign_finding_ids(merge_overlaps(findings))


def merge_overlaps(findings: List[Dict]) -> List[Dict]:
    ordered = sorted(findings, key=lambda item: (item["start"], -(item["end"] - item["start"])))
    kept = []
    for finding in ordered:
        if any(finding["start"] >= item["start"] and finding["end"] <= item["end"] for item in kept):
            continue
        kept.append(finding)
    return sorted(kept, key=lambda item: item["start"])


def assign_finding_ids(findings: List[Dict]) -> List[Dict]:
    for idx, finding in enumerate(findings):
        finding["id"] = idx
    return findings


def redact_text(text: str, findings: Iterable[Dict]) -> str:
    redacted = []
    cursor = 0
    for finding in sorted(findings, key=lambda item: item["start"]):
        start = finding["start"]
        end = finding["end"]
        if start < cursor:
            continue
        redacted.append(text[cursor:start])
        redacted.append(mask_span(text[start:end]))
        cursor = end
    redacted.append(text[cursor:])
    return "".join(redacted)


def redact_selected_text(text: str, findings: List[Dict], selected_ids: Iterable[int]) -> str:
    selected = {int(item) for item in selected_ids}
    return redact_text(text, [finding for finding in findings if finding["id"] in selected])


def mask_span(value: str, fill: str = MASK_CHAR) -> str:
    return "".join("\n" if char == "\n" else fill for char in value)


def public_findings(findings: Iterable[Dict]) -> List[Dict]:
    cleaned = []
    for finding in findings:
        item = dict(finding)
        item.pop("raw_value", None)
        cleaned.append(item)
    return cleaned


def calculate_risk(findings: Iterable[Dict]) -> str:
    scores = [RISK_SCORE.get(item.get("risk_level", "LOW"), 1) for item in findings]
    if not scores:
        return "LOW"
    max_score = max(scores)
    if max_score >= 4:
        return "CRITICAL"
    if max_score == 3:
        return "HIGH"
    if max_score == 2:
        return "MEDIUM"
    return "LOW"


def protected_output_path(session_dir: Path, filename: str, doc_id: str, file_type: str) -> Path:
    stem = secure_filename(Path(filename).stem) or "document"
    suffix = {"text": ".txt", "pdf": ".pdf", "docx": ".docx", "pptx": ".pptx"}[file_type]
    return session_dir / f"redacted_{doc_id}_{stem}{suffix}"


def save_redacted_file(doc: Dict, selected_ids: Iterable[int]) -> Path:
    original_path = Path(doc["original_path"])
    selected = {int(item) for item in selected_ids}
    findings = [finding for finding in doc["entities"] if finding["id"] in selected]
    output_path = protected_output_path(original_path.parent, doc["filename"], doc["id"], doc["file_type"])

    if doc["file_type"] == "text":
        output_path.write_text(redact_text(doc["extracted_text"], findings), encoding="utf-8")
    elif doc["file_type"] == "pdf":
        redact_pdf(original_path, output_path, findings)
    elif doc["file_type"] == "docx":
        redact_docx(original_path, output_path, findings)
    elif doc["file_type"] == "pptx":
        redact_pptx(original_path, output_path, findings)
    else:
        raise ValueError("Unsupported file type.")
    return output_path


def selected_raw_values(findings: List[Dict]) -> List[str]:
    values = []
    seen = set()
    for finding in sorted(findings, key=lambda item: item["length"], reverse=True):
        value = finding.get("raw_value", "")
        key = value.lower()
        if not value or key in seen:
            continue
        seen.add(key)
        values.append(value)
    return values


def redact_pdf(input_path: Path, output_path: Path, findings: List[Dict]) -> None:
    if fitz is None:
        raise RuntimeError("PDF visual redaction requires PyMuPDF. Run `pip install -r requirements.txt` and restart the app.")

    document = fitz.open(str(input_path))
    redaction_count = 0
    try:
        for finding in findings:
            value = finding.get("raw_value", "")
            if not value:
                continue
            page_number = finding.get("page")
            if page_number and 1 <= int(page_number) <= document.page_count:
                pages = [document[int(page_number) - 1]]
            else:
                pages = list(document)

            for page in pages:
                for rect in page.search_for(value):
                    padded = fitz.Rect(rect.x0 - 0.75, rect.y0 - 0.75, rect.x1 + 0.75, rect.y1 + 0.75)
                    page.add_redact_annot(padded, fill=(0, 0, 0))
                    redaction_count += 1

        if redaction_count == 0:
            raise ValueError("Selected text could not be located in the original PDF. The PDF may be scanned or use custom text encoding.")

        for page in document:
            page.apply_redactions()
        document.save(str(output_path), garbage=4, deflate=True, clean=True)
    finally:
        document.close()


def redact_docx(input_path: Path, output_path: Path, findings: List[Dict]) -> None:
    doc = DocxDocument(input_path)
    values = selected_raw_values(findings)
    for paragraph in doc.paragraphs:
        redact_docx_paragraph(paragraph, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    redact_docx_paragraph(paragraph, values)
    doc.save(output_path)


def find_value_matches(text: str, values: List[str]) -> List[Tuple[int, int]]:
    candidates = []
    for value in values:
        start = text.find(value)
        while start != -1:
            candidates.append((start, start + len(value)))
            start = text.find(value, start + len(value))
    candidates.sort(key=lambda item: (item[0], -(item[1] - item[0])))

    matches = []
    for start, end in candidates:
        if any(start < kept_end and end > kept_start for kept_start, kept_end in matches):
            continue
        matches.append((start, end))
    return matches


def docx_run_at(runs, position: int):
    cursor = 0
    for run in runs:
        next_cursor = cursor + len(run.text)
        if cursor <= position < next_cursor:
            return run
        cursor = next_cursor
    return runs[-1] if runs else None


def copy_docx_run_format(source, target) -> None:
    if source is None:
        return
    target.style = source.style
    target.bold = source.bold
    target.italic = source.italic
    target.underline = source.underline
    target.font.name = source.font.name
    target.font.size = source.font.size
    if source.font.highlight_color is not None:
        target.font.highlight_color = source.font.highlight_color
    try:
        if source.font.color.rgb is not None:
            target.font.color.rgb = source.font.color.rgb
    except (AttributeError, ValueError):
        pass


def apply_docx_mask_style(run) -> None:
    run.font.highlight_color = WD_COLOR_INDEX.BLACK
    run.font.color.rgb = DocxRGBColor(0, 0, 0)


def redact_docx_paragraph(paragraph, values: List[str]) -> None:
    text = paragraph.text
    matches = find_value_matches(text, values)
    if not matches:
        return

    paragraph_style = paragraph.style
    alignment = paragraph.alignment
    runs = list(paragraph.runs)
    paragraph.clear()
    paragraph.style = paragraph_style
    paragraph.alignment = alignment

    cursor = 0
    for start, end in matches:
        if cursor < start:
            source = docx_run_at(runs, cursor)
            run = paragraph.add_run(text[cursor:start])
            copy_docx_run_format(source, run)
        source = docx_run_at(runs, start)
        run = paragraph.add_run(mask_span(text[start:end]))
        copy_docx_run_format(source, run)
        apply_docx_mask_style(run)
        cursor = end
    if cursor < len(text):
        source = docx_run_at(runs, cursor)
        run = paragraph.add_run(text[cursor:])
        copy_docx_run_format(source, run)


def redact_pptx(input_path: Path, output_path: Path, findings: List[Dict]) -> None:
    prs = Presentation(input_path)
    values = selected_raw_values(findings)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame") or not shape.text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                redact_pptx_paragraph(paragraph, values)
    prs.save(output_path)


def pptx_run_at(runs, position: int):
    cursor = 0
    for run in runs:
        next_cursor = cursor + len(run.text)
        if cursor <= position < next_cursor:
            return run
        cursor = next_cursor
    return runs[-1] if runs else None


def copy_pptx_run_format(source, target) -> None:
    if source is None:
        return
    target.font.bold = source.font.bold
    target.font.italic = source.font.italic
    target.font.underline = source.font.underline
    target.font.name = source.font.name
    target.font.size = source.font.size
    try:
        if source.font.color.rgb is not None:
            target.font.color.rgb = source.font.color.rgb
    except (AttributeError, ValueError):
        pass


def apply_pptx_mask_style(run) -> None:
    run.font.color.rgb = PptxRGBColor(0, 0, 0)


def redact_pptx_paragraph(paragraph, values: List[str]) -> None:
    text = "".join(run.text for run in paragraph.runs)
    matches = find_value_matches(text, values)
    if not matches:
        return

    runs = list(paragraph.runs)
    paragraph.clear()
    cursor = 0
    for start, end in matches:
        if cursor < start:
            source = pptx_run_at(runs, cursor)
            run = paragraph.add_run()
            run.text = text[cursor:start]
            copy_pptx_run_format(source, run)
        source = pptx_run_at(runs, start)
        run = paragraph.add_run()
        run.text = mask_span(text[start:end])
        copy_pptx_run_format(source, run)
        apply_pptx_mask_style(run)
        cursor = end
    if cursor < len(text):
        source = pptx_run_at(runs, cursor)
        run = paragraph.add_run()
        run.text = text[cursor:]
        copy_pptx_run_format(source, run)


def summarize_documents(docs: Dict[str, Dict]) -> List[Dict]:
    rows = []
    for doc in docs.values():
        rows.append({
            "id": doc["id"],
            "filename": doc["filename"],
            "file_type": doc["file_type"],
            "uploaded_at": doc["uploaded_at"],
            "risk_level": doc["risk_level"],
            "entity_count": len(doc["entities"]),
            "has_protected": bool(doc.get("protected_path")) and Path(doc["protected_path"]).exists(),
        })
    return sorted(rows, key=lambda item: item["uploaded_at"], reverse=True)


@app.route("/")
@app.route("/dashboard")
def dashboard():
    return render_template("dashboard.html", initial_doc_id="")


@app.route("/findings/<doc_id>")
def findings_page(doc_id: str):
    return render_template("dashboard.html", initial_doc_id=doc_id)


@app.route("/api/documents")
def get_documents():
    sid = get_session_id()
    return jsonify(summarize_documents(SESSION_DOCS.get(sid, {})))


@app.route("/api/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400
    upload = request.files["file"]
    if not upload.filename:
        return jsonify({"error": "Invalid filename"}), 400

    sid = get_session_id()
    doc_id = str(uuid.uuid4())
    session_dir = SESSION_ROOT / sid
    filename = secure_filename(upload.filename)
    file_type = detect_file_type(filename)
    original_path = session_dir / f"original_{doc_id}_{filename}"
    upload.save(original_path)

    try:
        extracted_text, metadata = extract_text(original_path, file_type)
        if file_type == "pdf":
            findings = scan_pdf_pages(metadata.get("page_text", []))
        else:
            findings = scan_text(extracted_text)
        risk_level = calculate_risk(findings)
        public_metadata = dict(metadata)
        public_metadata.pop("page_text", None)

        record = {
            "id": doc_id,
            "filename": filename,
            "file_type": file_type,
            "uploaded_at": datetime.now().isoformat(),
            "original_path": str(original_path),
            "protected_path": None,
            "extracted_text": extracted_text,
            "entities": findings,
            "risk_level": risk_level,
            "extraction": {
                "char_count": len(extracted_text),
                "word_count": len(extracted_text.split()),
                "metadata": public_metadata,
            },
        }
        SESSION_DOCS[sid][doc_id] = record

        return jsonify({
            "success": True,
            "document_id": doc_id,
            "filename": filename,
            "file_type": file_type,
            "risk_level": risk_level,
            "entity_count": len(findings),
            "entities": public_findings(findings),
            "extraction": record["extraction"],
            "message": f"Scanned {len(extracted_text):,} characters and found {len(findings)} PII match(es).",
        })
    except Exception as exc:
        original_path.unlink(missing_ok=True)
        logger.exception("Upload processing failed")
        return jsonify({"error": f"Processing failed: {exc}"}), 500


@app.route("/api/report/<doc_id>")
def get_report(doc_id: str):
    sid = get_session_id()
    doc = SESSION_DOCS.get(sid, {}).get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found for this session"}), 404
    return jsonify({
        "document": doc["filename"],
        "uploaded": doc["uploaded_at"],
        "file_type": doc["file_type"],
        "risk_level": doc["risk_level"],
        "entities": public_findings(doc["entities"]),
        "entity_count": len(doc["entities"]),
        "extraction": doc["extraction"],
        "has_protected": bool(doc.get("protected_path")) and Path(doc["protected_path"]).exists(),
    })


@app.route("/api/protect/<doc_id>", methods=["POST"])
def protect_document(doc_id: str):
    sid = get_session_id()
    doc = SESSION_DOCS.get(sid, {}).get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found for this session"}), 404

    data = request.get_json() or {}
    selected_ids = data.get("selected_ids", [])
    valid_ids = {finding["id"] for finding in doc["entities"]}
    selected = []
    for item in selected_ids:
        try:
            finding_id = int(item)
        except (TypeError, ValueError):
            continue
        if finding_id in valid_ids:
            selected.append(finding_id)
    if not selected:
        return jsonify({"error": "Select at least one finding to redact"}), 400

    try:
        output_path = save_redacted_file(doc, selected)
        doc["protected_path"] = str(output_path)
        doc["redacted_ids"] = selected
        return jsonify({
            "success": True,
            "download_url": f"/api/download/{doc_id}/protected",
            "redacted_count": len(selected),
        })
    except Exception as exc:
        logger.exception("Redaction failed")
        return jsonify({"error": f"Redaction failed: {exc}"}), 500


@app.route("/api/download/<doc_id>/<version>")
def download(doc_id: str, version: str):
    sid = get_session_id()
    doc = SESSION_DOCS.get(sid, {}).get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found for this session"}), 404
    if version != "protected":
        return jsonify({"error": "Only redacted downloads are available"}), 400
    if not doc.get("protected_path"):
        return jsonify({"error": "Choose findings to redact before downloading"}), 400
    path = Path(doc["protected_path"])
    if not path.exists():
        return jsonify({"error": "Redacted file no longer exists"}), 404
    return send_file(path, as_attachment=True, download_name=f"redacted_{doc['filename']}")


@app.route("/api/delete/<doc_id>", methods=["DELETE"])
def delete_document(doc_id: str):
    sid = get_session_id()
    doc = SESSION_DOCS.get(sid, {}).pop(doc_id, None)
    if not doc:
        return jsonify({"error": "Document not found for this session"}), 404
    Path(doc["original_path"]).unlink(missing_ok=True)
    if doc.get("protected_path"):
        Path(doc["protected_path"]).unlink(missing_ok=True)
    return jsonify({"success": True})


@app.route("/api/session/end", methods=["POST"])
def end_session():
    cleanup_all_sessions()
    SESSION_ROOT.mkdir(exist_ok=True)
    session.clear()
    return jsonify({"success": True})


@app.errorhandler(413)
def too_large(_error):
    return jsonify({"error": "File too large. Maximum upload size is 100MB."}), 413


@app.errorhandler(500)
def server_error(_error):
    return jsonify({"error": "Internal server error"}), 500


if __name__ == "__main__":
    logger.info("Starting PrivacyGuard...")
    app.run(debug=False, port=5001, host="127.0.0.1")
